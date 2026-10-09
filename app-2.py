import os
import sys
import time
import string
import json
import base64
import threading
import webview
import pymupdf
import subprocess
import tempfile
import webbrowser
import textwrap
import copy
import shutil
import re
import urllib.request
import openai
from openai import OpenAI
try:
    import keyring
except Exception:
    keyring = None
from rank_bm25 import BM25Okapi
from ddgs import DDGS

try:
    import docx
except ImportError:
    docx = None

if getattr(sys, 'frozen', False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Private data (API keys, chats) lives in your user folder, NOT next to index.html,
# because the built-in web server serves everything that sits next to index.html.
DATA_DIR = os.path.join(os.path.expanduser("~"), "AIStudioPro")
os.makedirs(DATA_DIR, exist_ok=True)

DATA_FILENAMES = ["universal_config.json", "chat_history.json", "chat_settings.json", "chat_metadata.json", "prompts.json"]
for _name in DATA_FILENAMES:  # one-time migration of your existing data
    _old, _new = os.path.join(BASE_DIR, _name), os.path.join(DATA_DIR, _name)
    if os.path.exists(_old) and not os.path.exists(_new):
        try: shutil.copy2(_old, _new)
        except Exception: pass

CONFIG_FILE = os.path.join(DATA_DIR, "universal_config.json")
HISTORY_FILE = os.path.join(DATA_DIR, "chat_history.json")
SETTINGS_FILE = os.path.join(DATA_DIR, "chat_settings.json")
METADATA_FILE = os.path.join(DATA_DIR, "chat_metadata.json")
PROMPTS_FILE = os.path.join(DATA_DIR, "prompts.json")
INDEX_FILE = os.path.join(BASE_DIR, "index.html")

MAX_DOC_CHARS = 200000
MAX_WORKSPACE_FILE_BYTES = 1000000
BINARY_EXTS = {"zip","rar","7z","exe","dll","xlsx","xls","pptx","ppt","doc","mp3","mp4","mov","avi","mkv","wav","iso","bin","heic","tif","tiff","psd","svg"}
LOCAL_HOSTS = ("localhost", "127.0.0.1", "[::1]", "0.0.0.0")
KEYRING_SERVICE = "AIStudioPro"
BAD_TITLE_STARTS = ("here's a thinking", "here is a thinking", "thinking process", "okay, let", "okay let", "let me ", "the user ", "first, ", "<think")

VENDOR_DIR = os.path.join(BASE_DIR, "vendor")
VENDOR_FILES = {
    "tailwind.js": "https://cdn.tailwindcss.com",
    "marked.min.js": "https://cdn.jsdelivr.net/npm/marked/marked.min.js",
    "purify.min.js": "https://cdnjs.cloudflare.com/ajax/libs/dompurify/3.0.6/purify.min.js",
    "highlight.min.js": "https://cdnjs.cloudflare.com/ajax/libs/highlight.js/11.9.0/highlight.min.js",
    "github-dark.min.css": "https://cdnjs.cloudflare.com/ajax/libs/highlight.js/11.9.0/styles/github-dark.min.css",
}


def ensure_vendor_files():
    """Download the page libraries once into the 'vendor' folder so the app works offline afterwards.
    If a download is impossible right now, a tiny stand-in file loads the library from the internet instead."""
    try: os.makedirs(VENDOR_DIR, exist_ok=True)
    except Exception: return
    offline = False
    for name, url in VENDOR_FILES.items():
        path = os.path.join(VENDOR_DIR, name)
        try:
            if os.path.exists(path) and os.path.getsize(path) > 3000: continue   # already downloaded
        except OSError: pass
        ok = False
        if not offline:
            try:
                req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
                with urllib.request.urlopen(req, timeout=8) as r: data = r.read()
                head = data[:200].lstrip().lower()
                if len(data) > 3000 and not head.startswith(b"<!doctype html") and not head.startswith(b"<html"):
                    tmp = path + ".tmp"
                    with open(tmp, "wb") as f: f.write(data)
                    os.replace(tmp, path)
                    ok = True
            except Exception:
                offline = True    # no point trying the others right now
        if not ok:
            try:
                stub = (f'@import url("{url}");' if name.endswith(".css")
                        else f"document.write('<script src=\"{url}\"><\\/script>');")
                with open(path, "w", encoding="utf-8") as f: f.write(stub)
            except Exception: pass


def friendly_error(e):
    """Turn a technical API error into a sentence a normal person can act on."""
    A = openai
    if isinstance(e, A.AuthenticationError):
        msg = "Your API key was rejected. Open Settings, re-check the key for the active profile, and make sure it belongs to this provider."
    elif isinstance(e, A.PermissionDeniedError):
        msg = "Your key is not allowed to use this model or feature. Try another model, or check your provider account."
    elif isinstance(e, A.NotFoundError):
        msg = "The provider does not know this model or address. Check the Model ID and the Base URL in Settings."
    elif isinstance(e, A.RateLimitError):
        msg = "Too many requests, or your credit/quota ran out. Wait a minute, or check the billing page of your provider."
    elif isinstance(e, A.APITimeoutError):
        msg = "The provider took too long to answer (over 2 minutes). Try again, or pick a faster model."
    elif isinstance(e, A.APIConnectionError):
        msg = "Could not reach the provider. Check your internet connection and the Base URL (for LM Studio: make sure its local server is running)."
    elif isinstance(e, A.BadRequestError):
        msg = "The provider refused this request. Common causes: the conversation is too long, a file or image type the model cannot read, or a setting it does not support. Try a new chat or lower the Context Limit."
    elif isinstance(e, A.APIStatusError) and getattr(e, "status_code", 0) >= 500:
        msg = "The provider is having problems right now. Try again in a moment."
    else:
        msg = "Something went wrong while talking to the provider."
    detail = str(e).replace("`", "'")[:300]
    return f"**{msg}**\n\n*Technical details:* `{detail}`"


def js_toast(message):
    """Safely show a toast in the page (properly escaped)."""
    try:
        webview.windows[0].evaluate_js("showToast(" + json.dumps(str(message)) + ")")
    except Exception:
        pass

DEFAULT_PROFILES = {
    "active_profile_id": "default",
    "profiles": {
        "default": {"name": "OpenRouter", "base_url": "https://openrouter.ai/api/v1", "api_key": "", "model_name": "google/gemini-2.0-flash-001", "image_model": "dall-e-3", "custom_headers": ""},
        "local_gguf": {"name": "Local (LM Studio)", "base_url": "http://localhost:1234/v1", "api_key": "lm-studio", "model_name": "local-model", "image_model": "", "custom_headers": ""}
    }
}

class ChatEngine:
    def __init__(self):
        self.file_lock = threading.Lock()
        self.stream_lock = threading.Lock()
        self.stream_chunk = ""
        self.stream_is_done = False
        self.stream_error = None
        self.stream_gen = 0
        self.keyring_problem = ""
        self.keys_locked = False
        
        self.config = self._load_and_migrate_config()
        self.history = self._load_json(HISTORY_FILE, {})
        self._migrate_history_to_trees()
        self._repair_bad_titles()
        self.chat_settings = self._load_json(SETTINGS_FILE, {})
        self.metadata = self._load_json(METADATA_FILE, {})
        self.prompts = self._load_json(PROMPTS_FILE, [])
        
        for _cid in [c for c, d in self.history.items() if isinstance(d, dict) and not d.get("nodes")]:
            del self.history[_cid]  # drop empty leftover chats from older versions
        self.current_chat_id = f"chat_{int(time.time() * 1000)}"

        self.is_aborted = False
        self.active_stream = None
        self.workspace_name = None
        self.workspace_docs = []
        self.workspace_meta = []
        self.bm25_engine = None
        self.session_files = {}

    def poll_stream(self):
        with self.stream_lock:
            chunk = self.stream_chunk
            self.stream_chunk = ""
            done = self.stream_is_done
            error = self.stream_error
            if done:
                self.stream_error = None
                self.stream_is_done = False
            return {"chunk": chunk, "done": done, "error": error}

    def open_external_url(self, url):
        if url and (url.startswith("http://") or url.startswith("https://")):
            webbrowser.open(url)
            return True
        return False

    def save_code_to_file(self, code_string):
        try:
            dialog_type = getattr(webview.FileDialog, 'SAVE', getattr(webview, 'SAVE_DIALOG', 1))
            result = self._dialog_paths(webview.windows[0].create_file_dialog(dialog_type, save_filename='snippet.txt'))
            if not result: return None
            
            with open(result[0], 'w', encoding='utf-8') as f:
                f.write(code_string)
            js_toast("File saved successfully!")
            return True
        except Exception as e:
            js_toast(f"Save failed: {str(e)[:120]}")
            return False

    def export_chat_dialog(self, chat_id, export_format="md"):
        try:
            messages = self._public(self._get_linear_branch(self.history.get(chat_id)))
            if not messages: 
                js_toast("No messages to export.")
                return False
            
            content = ""
            if export_format == "json":
                content = json.dumps(messages, indent=2)
            else:
                content = f"# Chat Export: {messages[0].get('title', 'Conversation')}\n\n"
                for m in messages:
                    role_name = "**User**" if m['role'] == 'user' else "**Assistant**"
                    text = m.get('display_text', m.get('content', ''))
                    if isinstance(text, list): 
                        text = "\n".join([item.get('text', '') for item in text if item.get('type') == 'text'])
                    content += f"{role_name}:\n{text}\n\n---\n\n"
            
            dialog_type = getattr(webview.FileDialog, 'SAVE', getattr(webview, 'SAVE_DIALOG', 1))
            safe_title = "".join([c for c in messages[0].get('title', 'Conversation') if c.isalpha() or c.isdigit() or c==' ']).rstrip()
            safe_title = safe_title.replace(" ", "_")
            if not safe_title: safe_title = "Conversation"
            
            result = self._dialog_paths(webview.windows[0].create_file_dialog(
                dialog_type, 
                save_filename=f"{safe_title}.{export_format}"
            ))

            if result:
                with open(result[0], 'w', encoding='utf-8') as f:
                    f.write(content)
                js_toast(f"Success! Saved to: {result[0]}")
                return True
            else:
                js_toast("Export cancelled.")
                return False
        except Exception as e:
            js_toast(f"Export failed: {str(e)[:120]}")
            return False

    @staticmethod
    def _dialog_paths(result):
        """Different pywebview versions return a string or a list/tuple; always give a list."""
        if not result: return []
        if isinstance(result, str): return [result]
        return list(result)

    def _tokenize(self, text):
        # punctuation becomes a SPACE (self.user_name.get() -> self user name get) and camelCase is split
        text = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", text)
        text = re.sub(r"[^\w]+|_", " ", text)
        return text.lower().split()

    def _load_json(self, path, default):
        if os.path.exists(path):
            try:
                with open(path, "r", encoding="utf-8") as f: return json.load(f)
            except Exception:
                try: shutil.copy2(path, f"{path}.corrupt_{int(time.time())}.bak")  # never silently destroy a broken file
                except Exception: pass
        return default

    def _save_json(self, path, data):
        with self.file_lock:
            temp_path = f"{path}.tmp_{os.urandom(2).hex()}"
            try:
                text = None
                for _ in range(5):  # data may be changing in another thread; just retry
                    try:
                        text = json.dumps(data, indent=2, ensure_ascii=False)
                        break
                    except RuntimeError:
                        time.sleep(0.05)
                if text is None: raise RuntimeError("data kept changing while saving")
                with open(temp_path, "w", encoding="utf-8") as f:
                    f.write(text)
                os.replace(temp_path, path)
            except Exception as e:
                print(f"Error saving {path}: {e}")
                if os.path.exists(temp_path):
                    try: os.remove(temp_path)
                    except: pass

    # ---------- API key storage (Windows Credential Manager when available) ----------
    def _keyring_ok(self):
        if keyring is None:
            self.keyring_problem = "the 'keyring' package is not installed for this copy of Python"
            return False
        try:
            keyring.get_password(KEYRING_SERVICE, "__probe__")
            self.keyring_problem = ""
            return True
        except Exception as e:
            self.keyring_problem = f"Windows Credential Manager could not be reached: {str(e)[:80]}"
            return False

    def _load_and_migrate_config(self):
        data = self._load_json(CONFIG_FILE, None)
        if not data or "profiles" not in data:
            data = copy.deepcopy(DEFAULT_PROFILES)
            self.keys_in_keyring = False
            self.config = data
            self._write_config()
            return data
        use_keyring = self._keyring_ok()
        stored_ids = set(data.get("keyring_profiles", []))
        for pid, prof in data.get("profiles", {}).items():
            if use_keyring and pid in stored_ids and not prof.get("api_key"):
                try: prof["api_key"] = keyring.get_password(KEYRING_SERVICE, f"profile:{pid}") or ""
                except Exception: pass
        active = data.get("profiles", {}).get(data.get("active_profile_id", "default"), {})
        if active: data["api_key"] = active.get("api_key", "")
        self.keys_in_keyring = use_keyring
        self.keys_locked = bool(stored_ids) and not use_keyring   # keys exist in Credential Manager but this copy cannot read them
        self.config = data
        if use_keyring: self._write_config()    # moves any old plain-text keys into the secure store
        return data

    def _write_config(self):
        """Save settings. API keys go to the Windows Credential Manager when possible, otherwise into the file."""
        data = copy.deepcopy(self.config)
        previous = set(self.config.get("keyring_profiles", []))
        stored, use_keyring = [], self._keyring_ok()
        for pid, prof in data.get("profiles", {}).items():
            key = prof.get("api_key", "")
            if use_keyring:
                try:
                    if key: keyring.set_password(KEYRING_SERVICE, f"profile:{pid}", key)
                    else:
                        try: keyring.delete_password(KEYRING_SERVICE, f"profile:{pid}")
                        except Exception: pass
                    prof["api_key"] = ""
                    stored.append(pid)
                except Exception:
                    use_keyring = False            # could not store securely: keep the key in the file instead
        if use_keyring:
            data["api_key"] = ""
            for pid in previous - set(data.get("profiles", {}).keys()):   # deleted profiles
                try: keyring.delete_password(KEYRING_SERVICE, f"profile:{pid}")
                except Exception: pass
        if use_keyring:
            data["keyring_profiles"] = stored
        else:   # keep the marker for keys that still live in Credential Manager, so the .exe can find them again
            data["keyring_profiles"] = [pid for pid in previous if pid in data.get("profiles", {}) and not data["profiles"][pid].get("api_key")]
        self.keys_in_keyring = use_keyring
        self.config["keyring_profiles"] = data["keyring_profiles"]
        self._save_json(CONFIG_FILE, data)

    # ---------- chat management ----------
    def rename_chat(self, chat_id, title):
        title = re.sub(r"<[^>]*>", "", str(title or "")).strip()[:80]
        chat = self.history.get(chat_id)
        if not chat or not title: return False
        for node in chat.get("nodes", {}).values():
            if node.get("parent") is None: node["title"] = title
        self._save_json(HISTORY_FILE, self.history)
        return True

    def export_all_chats(self):
        try:
            payload = {"app": "AI Studio Pro", "version": 1, "history": self.history, "metadata": self.metadata, "settings": self.chat_settings}
            dialog_type = getattr(webview.FileDialog, 'SAVE', getattr(webview, 'SAVE_DIALOG', 1))
            result = self._dialog_paths(webview.windows[0].create_file_dialog(dialog_type, save_filename="ai_studio_backup.json"))
            if not result: return False
            with open(result[0], "w", encoding="utf-8") as f:
                json.dump(payload, f, ensure_ascii=False)
            js_toast(f"Backup saved: {result[0]}")
            return True
        except Exception as e:
            js_toast(f"Export failed: {str(e)[:120]}")
            return False

    def import_chats(self):
        """Adds chats from a backup file. Existing chats are never overwritten."""
        try:
            dialog_type = getattr(webview.FileDialog, 'OPEN', getattr(webview, 'OPEN_DIALOG', 0))
            result = self._dialog_paths(webview.windows[0].create_file_dialog(dialog_type, allow_multiple=False, file_types=("JSON files (*.json)", "All files (*.*)")))
            if not result: return 0
            with open(result[0], "r", encoding="utf-8") as f:
                data = json.load(f)
            incoming = data.get("history") if isinstance(data, dict) and "history" in data else data
            if not isinstance(incoming, dict): raise ValueError("This file is not an AI Studio Pro backup.")
            metadata = data.get("metadata", {}) if isinstance(data, dict) else {}
            settings = data.get("settings", {}) if isinstance(data, dict) else {}
            added = 0
            for cid, chat in incoming.items():
                if not (isinstance(chat, dict) and isinstance(chat.get("nodes"), dict) and chat["nodes"]): continue
                new_id = cid
                if new_id in self.history:
                    if json.dumps(self.history[new_id], sort_keys=True) == json.dumps(chat, sort_keys=True): continue   # already here
                    new_id = f"chat_{int(time.time() * 1000)}_{added}"
                self.history[new_id] = chat
                if isinstance(metadata.get(cid), dict): self.metadata[new_id] = metadata[cid]
                if isinstance(settings.get(cid), dict): self.chat_settings[new_id] = settings[cid]
                added += 1
            if added:
                self._save_json(HISTORY_FILE, self.history); self._save_json(METADATA_FILE, self.metadata); self._save_json(SETTINGS_FILE, self.chat_settings)
            js_toast(f"Imported {added} chat(s)." if added else "Nothing new to import.")
            return added
        except Exception as e:
            js_toast(f"Import failed: {str(e)[:120]}")
            return 0

    # ---------- images ----------
    @staticmethod
    def _image_bytes(src):
        """Returns (bytes, extension) for a data: URL or an http(s) image link."""
        if src.startswith("data:"):
            head, _, b64 = src.partition(",")
            data = base64.b64decode(b64)
        elif src.startswith(("http://", "https://")):
            req = urllib.request.Request(src, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=30) as r:
                data = r.read(40 * 1024 * 1024)
        else:
            raise ValueError("Unsupported image address.")
        ext = "png"
        if data[:3] == b"\xff\xd8\xff": ext = "jpg"
        elif data[:4] == b"RIFF" and data[8:12] == b"WEBP": ext = "webp"
        elif data[:4] == b"GIF8": ext = "gif"
        return data, ext

    def _auto_store_image(self, src):
        """Keeps a copy of every generated image on disk (provider links expire after a while)."""
        try:
            data, ext = self._image_bytes(src)
            folder = os.path.join(DATA_DIR, "images")
            os.makedirs(folder, exist_ok=True)
            path = os.path.join(folder, f"image_{int(time.time() * 1000)}.{ext}")
            with open(path, "wb") as f: f.write(data)
            js_toast(f"Image copy kept in {folder}")
        except Exception:
            pass

    def save_message_image(self, chat_id, msg_id):
        try:
            node = self.history.get(chat_id, {}).get("nodes", {}).get(msg_id)
            src = (node or {}).get("image")
            if not src: js_toast("No image found in this message."); return False
            data, ext = self._image_bytes(src)
            dialog_type = getattr(webview.FileDialog, 'SAVE', getattr(webview, 'SAVE_DIALOG', 1))
            result = self._dialog_paths(webview.windows[0].create_file_dialog(dialog_type, save_filename=f"image.{ext}"))
            if not result: return False
            with open(result[0], "wb") as f: f.write(data)
            js_toast(f"Image saved: {result[0]}")
            return True
        except Exception as e:
            js_toast(f"Could not save the image (the link may have expired): {str(e)[:100]}")
            return False

    def key_storage_info(self):
        if getattr(self, "keys_in_keyring", False):
            return {"secure": True, "text": "🔒 Your API keys are stored in Windows Credential Manager (not in a file)."}
        why = f" ({self.keyring_problem})" if self.keyring_problem else ""
        if self.keys_locked:
            return {"secure": False, "text": f"⚠️ Your keys are stored in Windows Credential Manager, but this copy of the app cannot reach it{why}. Run the .exe, or install it with: pip install keyring (in the same Python that VS Code uses)."}
        return {"secure": False, "text": f"⚠️ Secure storage is not available{why} - keys are saved in a file in your AIStudioPro folder."}

    def _migrate_history_to_trees(self):
        dirty = False
        for cid, chat_data in list(self.history.items()):
            if isinstance(chat_data, list):
                dirty = True
                new_chat = {"nodes": {}, "current_node": None}
                prev_id = None
                for m in chat_data:
                    nid = "msg_" + os.urandom(4).hex()
                    m["id"] = nid
                    m["parent"] = prev_id
                    m["children"] = []
                    new_chat["nodes"][nid] = m
                    if prev_id: new_chat["nodes"][prev_id]["children"].append(nid)
                    prev_id = nid
                new_chat["current_node"] = prev_id
                self.history[cid] = new_chat
        if dirty: self._save_json(HISTORY_FILE, self.history)

    def _get_linear_branch(self, chat_data):
        if not chat_data or not chat_data.get("current_node"): return []
        nodes = chat_data["nodes"]
        curr = chat_data["current_node"]
        path = []
        while curr and curr in nodes:
            path.append(curr)
            curr = nodes[curr]["parent"]
        path.reverse()
        linear = []
        for nid in path:
            msg = dict(nodes[nid])
            pid = msg["parent"]
            if pid and pid in nodes:
                siblings = nodes[pid]["children"]
                msg["branch_count"] = len(siblings)
                msg["branch_index"] = siblings.index(nid) + 1
            else:
                msg["branch_count"] = 1
                msg["branch_index"] = 1
            linear.append(msg)
        return linear

    def _add_node(self, chat_id, role, content, display_text, title, is_image_mode, files=None, image=None):
        chat = self.history[chat_id]
        nid = "msg_" + os.urandom(4).hex()
        pid = chat.get("current_node")
        msg = {
            "id": nid, "parent": pid, "children": [], "role": role, 
            "content": content, "display_text": display_text, "title": title, "is_image_mode": is_image_mode,
            "files": files or [], "image": image
        }
        chat["nodes"][nid] = msg
        if pid and pid in chat["nodes"]: chat["nodes"][pid]["children"].append(nid)
        chat["current_node"] = nid
        return nid

    def switch_branch(self, chat_id, msg_id, direction):
        chat = self.history.get(chat_id)
        if not chat or msg_id not in chat["nodes"]: return False
        msg = chat["nodes"][msg_id]
        pid = msg["parent"]
        if not pid: return False
        siblings = chat["nodes"][pid]["children"]
        idx = siblings.index(msg_id)
        new_idx = (idx + direction) % len(siblings)
        curr = siblings[new_idx]
        while chat["nodes"][curr]["children"]:
            curr = chat["nodes"][curr]["children"][-1]
        chat["current_node"] = curr
        self._save_json(HISTORY_FILE, self.history)
        return True

    def get_config(self): return self.config

    def list_models(self):
        """Ask the active provider which models it offers (used by the model picker)."""
        try:
            active_id = self.config.get("active_profile_id", "default")
            profile = self.config.get("profiles", {}).get(active_id, {})
            api_key, base_url = profile.get("api_key") or "", profile.get("base_url") or ""
            if not api_key and not any(h in base_url for h in LOCAL_HOSTS):
                return {"ok": False, "error": "no API key yet", "models": []}
            raw = (profile.get("custom_headers") or "").strip()
            headers = None
            if raw.startswith("{"):
                try: headers = json.loads(raw)
                except Exception: pass
            client = OpenAI(base_url=base_url or None, api_key=api_key or "not-needed", default_headers=headers, timeout=15.0)
            seen, models = set(), []
            for m in client.models.list():
                mid = getattr(m, "id", None)
                if not mid or mid in seen: continue
                seen.add(mid)
                name = getattr(m, "name", None) or (getattr(m, "model_extra", None) or {}).get("name") or mid
                models.append({"id": str(mid), "name": str(name)})
                if len(models) >= 2000: break
            models.sort(key=lambda x: x["name"].lower())
            return {"ok": True, "models": models}
        except Exception as e:
            return {"ok": False, "error": str(e)[:150], "models": []}
    @staticmethod
    def _public(messages):
        return [{k: v for k, v in m.items() if k not in ("llm_text", "ctx_sources")} for m in messages]

    def get_history(self): return {cid: self._public(self._get_linear_branch(cdata)) for cid, cdata in list(self.history.items())}
    def get_metadata(self): return self.metadata
    def get_chat_settings(self, chat_id): return self.chat_settings.get(chat_id, {})
    def get_prompts(self): return self.prompts

    def save_prompt(self, title, content, prompt_id=None):
        if prompt_id:
            for p in self.prompts:
                if p["id"] == prompt_id:
                    p["title"], p["content"] = title, content
                    break
        else:
            self.prompts.append({"id": f"p_{os.urandom(4).hex()}", "title": title, "content": content})
        self._save_json(PROMPTS_FILE, self.prompts)
        return self.prompts

    def delete_prompt(self, prompt_id):
        self.prompts = [p for p in self.prompts if p["id"] != prompt_id]
        self._save_json(PROMPTS_FILE, self.prompts)
        return self.prompts

    def save_config(self, cfg):
        self.config.update(cfg)
        active_id = self.config.get("active_profile_id", "default")
        profile = self.config.get("profiles", {}).get(active_id, {})
        if profile:
            self.config["api_key"] = profile.get("api_key", "")
            self.config["base_url"] = profile.get("base_url", "")
            self.config["model_name"] = profile.get("model_name", "")
            self.config["image_model"] = profile.get("image_model", "dall-e-3")
        self._write_config()
        return True

    def start_new_chat(self, is_image_mode=False):
        self.current_chat_id = f"chat_{int(time.time() * 1000)}"
        self.session_files.clear()
        return self.current_chat_id

    def set_current_chat(self, chat_id): 
        self.current_chat_id = chat_id

    def toggle_pin_chat(self, chat_id):
        if chat_id not in self.metadata: self.metadata[chat_id] = {}
        self.metadata[chat_id]["pinned"] = not self.metadata[chat_id].get("pinned", False)
        self._save_json(METADATA_FILE, self.metadata)
        return self.metadata[chat_id]["pinned"]

    def delete_chat(self, chat_id):
        for store, file_path in [(self.history, HISTORY_FILE), (self.chat_settings, SETTINGS_FILE), (self.metadata, METADATA_FILE)]:
            if chat_id in store:
                del store[chat_id]
                self._save_json(file_path, store)
        return True

    def clear_all_chats(self):
        self.history, self.chat_settings, self.metadata = {}, {}, {}
        self._save_json(HISTORY_FILE, self.history); self._save_json(SETTINGS_FILE, self.chat_settings); self._save_json(METADATA_FILE, self.metadata)
        self.session_files.clear()
        return True

    def stop_generation(self):
        self.is_aborted = True
        if self.active_stream:
            try: self.active_stream.close()
            except: pass

    def rewind_chat(self, chat_id, msg_id):
        chat = self.history.get(chat_id)
        if chat and msg_id in chat["nodes"]:
            chat["current_node"] = chat["nodes"][msg_id]["parent"]
            self._save_json(HISTORY_FILE, self.history)
        return True

    def _get_python_interpreter(self):
        if not getattr(sys, 'frozen', False):
            return sys.executable
        for cmd in ['python', 'py', 'python3']:
            found = shutil.which(cmd)
            if found and "windowsapps" not in found.lower():  # skip the Microsoft Store stub
                if os.path.abspath(found).lower() != os.path.abspath(sys.executable).lower():
                    return found
        possible_dirs = [
            os.path.expandvars(r"%LOCALAPPDATA%\Programs\Python"),
            r"C:\Python312", r"C:\Python311", r"C:\Python310", r"C:\Python39"
        ]
        for base in possible_dirs:
            if os.path.exists(base):
                for root, dirs, files in os.walk(base):
                    if "python.exe" in files:
                        p_path = os.path.join(root, "python.exe")
                        if "scripts" not in p_path.lower():
                            return p_path
        return None

    def execute_local_code(self, code_string):
        python_bin = self._get_python_interpreter()
        if not python_bin:
            return "Execution Error: Python is not installed or not in your Windows PATH. Please install Python to run scripts locally."
        
        temp_path = None
        try:
            with tempfile.NamedTemporaryFile(mode='w', suffix='.py', delete=False, encoding='utf-8') as f:
                f.write(code_string)
                temp_path = f.name
            
            creation_flags = 0
            if sys.platform == 'win32':
                creation_flags = subprocess.CREATE_NO_WINDOW

            env = dict(os.environ, PYTHONUTF8="1", PYTHONIOENCODING="utf-8")
            result = subprocess.run(
                [python_bin, temp_path],
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                stdin=subprocess.DEVNULL,
                env=env,
                timeout=15,
                creationflags=creation_flags
            )
            output = result.stdout
            if result.stderr:
                output += f"\nErrors:\n{result.stderr}"
            output = output.strip()
            if len(output) > 20000: output = output[:20000] + "\n...[output cut]"
            return output or "Code executed successfully with no console output."
        except subprocess.TimeoutExpired:
            return "Execution Error: Script timed out after 15 seconds."
        except Exception as e:
            return f"Execution Error: {str(e)}"
        finally:
            if temp_path and os.path.exists(temp_path):
                try: os.remove(temp_path)
                except: pass

    def attach_workspace(self):
        try:
            dialog_type = getattr(webview.FileDialog, 'FOLDER', getattr(webview, 'FOLDER_DIALOG', 2))
            result = self._dialog_paths(webview.windows[0].create_file_dialog(dialog_type))
            if not result: return None
            folder_path = result[0]
            js_toast(f"Scanning {os.path.basename(folder_path)}...")
            threading.Thread(target=self._index_folder, args=(folder_path,), daemon=True).start()
            return os.path.basename(folder_path)
        except Exception:
            return None

    def _index_folder(self, folder_path):
        try:
            self.bm25_engine = None
            name = os.path.basename(folder_path)
            exclude_dirs = {'.gradle', '.idea', 'build', 'node_modules', '.git', 'intermediates', '__pycache__', '.venv', 'venv', 'dist', '.next', 'vendor'}
            text_exts = {'kt', 'kts', 'py', 'md', 'txt', 'java', 'xml', 'json', 'csv', 'html', 'css', 'js', 'ts', 'tsx', 'jsx',
                         'c', 'h', 'cpp', 'hpp', 'cs', 'go', 'rs', 'rb', 'php', 'swift', 'sh', 'bat', 'ps1', 'sql',
                         'yml', 'yaml', 'toml', 'ini', 'cfg', 'gradle', 'properties', 'rst', 'tex'}
            docs, meta = [], []   # built locally and swapped in at the end, so searches never see half-built data

            for root, dirs, files in os.walk(folder_path):
                dirs[:] = [d for d in dirs if d not in exclude_dirs]
                for file in files:
                    ext = file.rsplit('.', 1)[-1].lower() if '.' in file else ''
                    path = os.path.join(root, file)
                    rel = os.path.relpath(path, folder_path)
                    text = ""
                    if file.lower().endswith(('.min.js', '.min.css', '.map')): continue   # minified libraries are just noise
                    try:
                        if os.path.getsize(path) > MAX_WORKSPACE_FILE_BYTES and ext != 'pdf': continue
                    except OSError: continue
                    if ext in text_exts:
                        try:
                            with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                                text = f.read()
                        except Exception: pass
                    elif ext == 'pdf':
                        try:
                            with pymupdf.open(path) as pdf:
                                text = "\n".join(page.get_text() for page in pdf)
                        except Exception: pass

                    if text:
                        buf, size = [], 0
                        for line in text.splitlines():
                            line = line[:1000]
                            if size + len(line) > 1200 and buf:
                                docs.append("\n".join(buf)); meta.append({"filename": rel})
                                buf, size = buf[-3:], sum(len(x) for x in buf[-3:])
                            buf.append(line); size += len(line) + 1
                        if buf and "".join(buf).strip():
                            docs.append("\n".join(buf)); meta.append({"filename": rel})

            if docs:
                engine = BM25Okapi([self._tokenize(d) or ["_"] for d in docs])
                self.workspace_docs, self.workspace_meta, self.workspace_name = docs, meta, name
                self.bm25_engine = engine
                js_toast(f"Indexed {len(docs)} chunks successfully.")
            else:
                self.workspace_docs, self.workspace_meta, self.workspace_name = [], [], name
                js_toast("No readable text files found.")
        except Exception as e:
            js_toast(f"Indexing error: {str(e)[:120]}")

    def _process_frontend_file(self, file_obj):
        try:
            name, data_url = file_obj.get("name", "unnamed"), file_obj.get("data", "")
            ext = name.rsplit('.', 1)[-1].lower() if '.' in name else ''
            if not data_url or "," not in data_url: return None
            header, b64_data = data_url.split(",", 1)

            if ext in ['png', 'jpg', 'jpeg', 'webp', 'gif'] or header.startswith("data:image/") and ext not in BINARY_EXTS:
                return {"name": name, "is_image": True, "data_url": data_url}
            raw_bytes = base64.b64decode(b64_data)
            text = None
            if ext == 'pdf':
                with pymupdf.open(stream=raw_bytes, filetype="pdf") as pdf:
                    text = "\n".join(p.get_text() for p in pdf)
                if not text.strip(): text = "[This PDF has no readable text - it may be a scanned image.]"
            elif ext == 'docx':
                if not docx: text = "[Cannot read .docx: run  pip install python-docx  and restart the app.]"
                else:
                    import io
                    d = docx.Document(io.BytesIO(raw_bytes))
                    parts = [p.text for p in d.paragraphs]
                    for t in d.tables:
                        for row in t.rows: parts.append(" | ".join(c.text for c in row.cells))
                    text = "\n".join(parts)
            elif ext in BINARY_EXTS or b"\x00" in raw_bytes[:4096]:
                text = f"[File '{name}' is a binary file type and cannot be read as text.]"
            else:
                text = raw_bytes.decode('utf-8', errors='ignore')
            if len(text) > MAX_DOC_CHARS: text = text[:MAX_DOC_CHARS] + "\n...[document cut: too long]"
            return {"name": name, "is_image": False, "text": text}
        except Exception as e:
            return {"name": file_obj.get("name", "file"), "is_image": False, "text": f"[Could not read this file: {e}]"}

    @staticmethod
    def _clean_title(raw, prompt):
        """Short, clean chat title. Falls back to the first words of the user's message if the model returns junk."""
        fallback = re.sub(r"\s+", " ", str(prompt or "Conversation")).strip()[:28] or "Conversation"
        t = re.sub(r"<think>.*?</think>", "", str(raw or ""), flags=re.S | re.I)
        t = re.sub(r"<[^>]*>", "", t)
        t = t.strip().splitlines()[0].strip() if t.strip() else ""
        t = re.sub(r"[*#_`>\"]", "", t).strip(" .:-")
        if not t or len(t.split()) > 8 or len(t) > 60 or t.lower().startswith(BAD_TITLE_STARTS):
            return fallback
        return t

    def _repair_bad_titles(self):
        """Fix titles saved by older versions when a reasoning model wrote its thoughts as the title."""
        changed = False
        for chat in self.history.values():
            nodes = chat.get("nodes", {}) if isinstance(chat, dict) else {}
            for node in nodes.values():
                if node.get("parent") is None and node.get("role") == "user":
                    title = str(node.get("title", ""))
                    if title and (len(title) > 60 or title.lower().startswith(BAD_TITLE_STARTS) or "**" in title):
                        node["title"] = self._clean_title(title, node.get("display_text") or "Conversation")
                        changed = True
        if changed: self._save_json(HISTORY_FILE, self.history)

    def _generate_title_thread(self, chat_id, prompt_text, profile):
        try:
            custom_headers_raw = (profile.get("custom_headers") or "").strip()
            headers = json.loads(custom_headers_raw) if custom_headers_raw.startswith("{") else None
            client = OpenAI(base_url=profile.get("base_url") or None, api_key=profile.get("api_key") or "none", default_headers=headers, timeout=60.0)
            response = client.chat.completions.create(
                model=profile.get("model_name") or "gpt-4o",
                messages=[{"role": "system", "content": "Summarize this user request in 2 to 4 words. Respond ONLY with the title. No quotes, no punctuation."}, {"role": "user", "content": str(prompt_text)[:500]}],
                max_tokens=40, temperature=0.3
            )
            title = self._clean_title(response.choices[0].message.content, prompt_text)
            chat = self.history.get(chat_id)
            if chat and chat.get("current_node"):
                curr = chat["current_node"]
                while chat["nodes"][curr]["parent"]: curr = chat["nodes"][curr]["parent"]
                chat["nodes"][curr]["title"] = title
                self._save_json(HISTORY_FILE, self.history)
        except Exception: pass

    # ---------- chat engine ----------
    def _emit(self, gen, chunk=None, error=None, done=False):
        """Send text/errors/'finished' to the page, but only if this worker is still the current one."""
        with self.stream_lock:
            if gen != self.stream_gen: return
            if chunk: self.stream_chunk += chunk
            if error: self.stream_error = error
            if done: self.stream_is_done = True

    def _add_error_node(self, chat_id, text):
        nid = self._add_node(chat_id, "assistant", text, text, "", False)
        self.history[chat_id]["nodes"][nid]["is_error"] = True
        return nid

    def _turn_content(self, m, limit=None):
        """Content to send to the model for a stored message (includes attached document text)."""
        content = m.get("content", "")
        llm_text = m.get("llm_text")
        if llm_text:
            if isinstance(content, list):
                content = copy.deepcopy(content)
                for item in content:
                    if item.get("type") == "text":
                        item["text"] = llm_text
                        break
                else:
                    content.insert(0, {"type": "text", "text": llm_text})
            else:
                content = llm_text
        if limit:
            if isinstance(content, str) and len(content) > limit:
                content = content[:limit] + "\n...[older message shortened]..."
            elif isinstance(content, list):
                content = copy.deepcopy(content)
                for item in content:
                    if item.get("type") == "text" and isinstance(item.get("text"), str) and len(item["text"]) > limit:
                        item["text"] = item["text"][:limit] + "\n...[older message shortened]..."
        return content

    def send_prompt_stream(self, user_prompt, file_data_list, params):
        with self.stream_lock:
            self.is_aborted = False
            self.stream_gen += 1
            gen = self.stream_gen
            self.stream_chunk = ""
            self.stream_is_done = False
            self.stream_error = None
        threading.Thread(target=self._stream_worker, args=(gen, user_prompt, file_data_list, params), daemon=True).start()

    def _stream_worker(self, gen, user_prompt, file_data_list, params):
        sources_list = []
        target_chat = self.current_chat_id
        if target_chat not in self.history:
            self.history[target_chat] = {"nodes": {}, "current_node": None}

        self.chat_settings[target_chat] = params
        self._save_json(SETTINGS_FILE, self.chat_settings)

        active_id = self.config.get("active_profile_id", "default")
        profile = self.config.get("profiles", {}).get(active_id, {})
        if profile:   # use ONLY the chosen profile (never borrow another profile's key)
            api_key = profile.get("api_key") or ""
            base_url = profile.get("base_url") or ""
            model_name = profile.get("model_name") or "google/gemini-2.0-flash-001"
        else:
            api_key = self.config.get("api_key", "")
            base_url = self.config.get("base_url", "")
            model_name = self.config.get("model_name") or "google/gemini-2.0-flash-001"

        custom_headers_raw = (profile.get("custom_headers") or "").strip()
        custom_headers = None
        if custom_headers_raw.startswith("{") and custom_headers_raw.endswith("}"):
            try: custom_headers = json.loads(custom_headers_raw)
            except Exception: self._emit(gen, chunk="> ⚠️ *Custom headers are not valid JSON and were ignored.*\n\n")

        is_image_gen = params.get("is_image_mode", False)
        is_regenerate = params.get("regenerate", False)
        active_workspace = params.get("active_workspace")

        is_local = any(h in str(base_url) for h in LOCAL_HOSTS)
        if not api_key and not is_local:
            if not is_regenerate:
                self._add_node(target_chat, "user", user_prompt, user_prompt, (user_prompt or "Chat")[:28], False)
            err_msg = "**Configuration Warning**: No API Key detected. Please open Settings and configure your Provider Profile."
            self._add_error_node(target_chat, err_msg)
            self._save_json(HISTORY_FILE, self.history)
            self._emit(gen, error=err_msg, done=True)
            return
        if not api_key: api_key = "not-needed"

        current_turn_payload = None
        workspace_sources = []
        if not is_regenerate:
            doc_texts, image_urls, doc_meta_list = [], [], []

            if file_data_list:
                for f in file_data_list:
                    processed = self._process_frontend_file(f)
                    if not processed: continue
                    if processed["is_image"]: image_urls.append(processed["data_url"])
                    else:
                        doc_texts.append(f"--- Document: {processed['name']} ---\n{processed['text']}")
                        doc_meta_list.append({"name": processed["name"]})

            workspace_context = ""
            if active_workspace and user_prompt and user_prompt.strip():
                engine, docs, metas = self.bm25_engine, self.workspace_docs, self.workspace_meta   # snapshot
                if not engine:
                    self._emit(gen, chunk="> ⚠️ *The workspace folder is still being scanned or has no readable files.*\n\n")
                else:
                    try:
                        scores = engine.get_scores(self._tokenize(user_prompt))
                        top = [i for i in sorted(range(len(scores)), key=lambda i: scores[i], reverse=True)[:6] if scores[i] > 0 and i < len(docs)]
                        if top:
                            workspace_context += f"\n[Local Workspace Files: {active_workspace}]\n"
                            for idx in top:
                                workspace_context += f"--- File: {metas[idx]['filename']} ---\n{docs[idx]}\n\n"
                                fname = metas[idx]["filename"]
                                existing = next((s for s in workspace_sources if s["file"] == fname), None)
                                if existing:
                                    existing["chunks"] += 1
                                    existing["snippet"] += "\n\n[...]\n\n" + docs[idx][:400]
                                else:
                                    workspace_sources.append({"file": fname, "chunks": 1, "snippet": docs[idx][:600]})
                        else:
                            self._emit(gen, chunk="> ℹ️ *No matching text found in the workspace for this question.*\n\n")
                    except Exception: pass

            is_web_search = params.get("web_search", False)
            web_context = ""
            if is_web_search and user_prompt and user_prompt.strip():
                try:
                    self._emit(gen, chunk="> 🌐 *Searching live web...*\n\n")
                    query = " ".join(user_prompt.split())[:300]
                    results = list(DDGS().text(query, max_results=4))
                    if results:
                        web_context += "\n[Live Web Search Reference Material]\n"
                        for res in results:
                            title = str(res.get('title', 'Web Result'))
                            url = str(res.get('href', ''))
                            body = str(res.get('body', ''))
                            if url.startswith("http://") or url.startswith("https://"):
                                sources_list.append({"title": title, "url": url})
                                web_context += f"Source: {title} ({url})\nSnippet: {body}\n\n"
                    else:
                        web_context += "\n[Live Web Search: No matching results found online.]\n"
                except Exception as e:
                    self._emit(gen, chunk=f"> ⚠️ *Web search failed ({str(e)[:100]}). Answering without it.*\n\n")

            prompt_for_model = user_prompt if (user_prompt and user_prompt.strip()) else ("(see the attached file)" if (doc_texts or image_urls) else "")
            has_extra = bool(doc_texts or workspace_context or web_context)
            combined_text = ""
            if has_extra:
                if workspace_context: combined_text += f"{workspace_context}\n"
                if doc_texts: combined_text += "\n\n".join(doc_texts) + "\n\n"
                if web_context: combined_text += f"{web_context}\n"
                combined_text += f"User Request: {prompt_for_model}"

            if image_urls or has_extra:
                current_turn_payload = [{"type": "text", "text": combined_text or prompt_for_model}]
                for url in image_urls: current_turn_payload.append({"type": "image_url", "image_url": {"url": url}})
            else:
                current_turn_payload = prompt_for_model

            if image_urls:
                clean_history_content = [{"type": "text", "text": prompt_for_model}]
                for url in image_urls: clean_history_content.append({"type": "image_url", "image_url": {"url": url}})
            else:
                clean_history_content = prompt_for_model

            display_title = (user_prompt or "Attachment")[:28]
            display_label = user_prompt or f"[{len(file_data_list or [])} Files Attached]"
            primary_image = image_urls[0] if image_urls else None

            nid = self._add_node(target_chat, "user", clean_history_content, display_label, display_title, is_image_gen, files=doc_meta_list, image=primary_image)
            if has_extra:   # remember the attached text so follow-ups and Regenerate still see it
                self.history[target_chat]["nodes"][nid]["llm_text"] = combined_text
                self.history[target_chat]["nodes"][nid]["ctx_sources"] = {"web": sources_list, "workspace": workspace_sources}

            if len(self._get_linear_branch(self.history[target_chat])) == 1 and not is_image_gen:
                threading.Thread(target=self._generate_title_thread, args=(target_chat, user_prompt or "Attachment", {"api_key": api_key, "base_url": base_url, "model_name": model_name, "custom_headers": profile.get("custom_headers", "")}), daemon=True).start()

        linear_history = self._get_linear_branch(self.history[target_chat])
        if is_regenerate and linear_history:      # keep the same sources when an answer is regenerated
            ctx = linear_history[-1].get("ctx_sources") or {}
            sources_list = list(ctx.get("web", []))
            workspace_sources = list(ctx.get("workspace", []))
        client = OpenAI(base_url=base_url or None, api_key=api_key, default_headers=custom_headers, timeout=120.0)
        assistant_node_id = None
        full_reply = ""

        if is_image_gen:
            try:
                img_prompt = linear_history[-1].get("display_text", "") if linear_history else ""
                image_model = profile.get("image_model") or "dall-e-3"
                response = client.images.generate(model=image_model, prompt=img_prompt, size="1024x1024", n=1)
                item = response.data[0]
                img_url = getattr(item, "url", None)
                if not img_url and getattr(item, "b64_json", None):
                    img_url = "data:image/png;base64," + item.b64_json
                if not img_url: raise RuntimeError("The provider returned no image. Check that your provider supports image generation.")
                output_md = f"![Generated Image]({img_url})"
                if not img_url.startswith("data:"): output_md += f"\n\n[Open Full Size Image]({img_url})"
                nid = self._add_node(target_chat, "assistant", output_md, output_md, "", False, image=img_url)
                self.history[target_chat]["nodes"][nid]["model"] = profile.get("image_model") or "dall-e-3"
                self._save_json(HISTORY_FILE, self.history)
                self._emit(gen, chunk=output_md, done=True)
                threading.Thread(target=self._auto_store_image, args=(img_url,), daemon=True).start()
            except Exception as e:
                err = "**Image generation failed.** Not every provider supports images - check yours.\n\n" + friendly_error(e)
                self._add_error_node(target_chat, err)
                self._save_json(HISTORY_FILE, self.history)
                self._emit(gen, error=err, done=True)
            return

        try:
            limit = params.get("context_limit", 10)
            messages_payload = [{"role": "system", "content": params["system_prompt"]}] if params.get("system_prompt") else []

            earlier = []
            for m in linear_history[:-1]:      # leave out failed turns so errors never confuse the model
                if m.get("is_error"):
                    if earlier and earlier[-1]["role"] == "user": earlier.pop()
                    continue
                earlier.append(m)
            for m in earlier[-(limit - 1):] if limit > 1 else []:
                messages_payload.append({"role": m["role"], "content": self._turn_content(m, 8000)})

            if linear_history:
                last_turn = linear_history[-1]
                active_content = current_turn_payload if (current_turn_payload and not is_regenerate) else self._turn_content(last_turn)
                messages_payload.append({"role": last_turn["role"], "content": active_content})

            kwargs = dict(model=model_name, messages=messages_payload, stream=True,
                          temperature=params.get("temperature", 0.7), stream_options={"include_usage": True})
            while True:   # drop settings that this provider/model refuses, then try again
                try:
                    self.active_stream = client.chat.completions.create(**kwargs)
                    break
                except Exception as setup_err:
                    text = str(setup_err).lower()
                    if "stream_options" in kwargs and ("stream_options" in text or "include_usage" in text or isinstance(setup_err, openai.BadRequestError)):
                        kwargs.pop("stream_options")
                    elif "temperature" in kwargs and "temperature" in text:
                        kwargs.pop("temperature")
                    else:
                        raise

            assistant_node_id = self._add_node(target_chat, "assistant", "", "", "", False)

            if active_workspace and not is_regenerate and self.bm25_engine:
                status_msg = f"> 📂 *Scanned local workspace ({active_workspace}).*\n\n"
                full_reply += status_msg
                self._emit(gen, chunk=status_msg)

            model_text = ""
            usage = None
            cancel_msg = "\n\n*(Cancelled by user)*"
            try:
                for chunk in self.active_stream:
                    if self.is_aborted or gen != self.stream_gen:
                        try: self.active_stream.close()
                        except Exception: pass
                        full_reply += cancel_msg
                        self._emit(gen, chunk=cancel_msg)
                        break
                    if getattr(chunk, "usage", None): usage = chunk.usage
                    if chunk.choices and chunk.choices[0].delta.content:
                        token = chunk.choices[0].delta.content
                        model_text += token
                        full_reply += token
                        self._emit(gen, chunk=token)
            except Exception:
                if not (self.is_aborted or gen != self.stream_gen): raise
                if cancel_msg not in full_reply:        # stream was closed by the Stop button: keep what we have
                    full_reply += cancel_msg
                    self._emit(gen, chunk=cancel_msg)

            if not model_text.strip() and cancel_msg not in full_reply:
                note = "*(The model returned an empty reply. Try again or pick another model.)*"
                full_reply += note
                self._emit(gen, chunk=note)

            node = self.history[target_chat]["nodes"][assistant_node_id]
            node["content"] = full_reply
            node["display_text"] = full_reply
            node["sources"] = sources_list
            node["workspace_sources"] = workspace_sources
            node["model"] = model_name
            if usage is not None:
                node["usage"] = {"prompt": getattr(usage, "prompt_tokens", None), "completion": getattr(usage, "completion_tokens", None)}
            self._save_json(HISTORY_FILE, self.history)
            self._emit(gen, done=True)

        except Exception as e:
            err = "\n\n" + friendly_error(e)
            if assistant_node_id:
                node = self.history[target_chat]["nodes"][assistant_node_id]
                node["content"] = full_reply + err
                node["display_text"] = full_reply + err
                node["is_error"] = True
            else:
                self._add_error_node(target_chat, err)
            self._save_json(HISTORY_FILE, self.history)
            self._emit(gen, error=err, done=True)


if __name__ == '__main__':
    import multiprocessing
    multiprocessing.freeze_support()
    app_cache_dir = os.path.join(DATA_DIR, 'webview_cache')
    os.makedirs(app_cache_dir, exist_ok=True)
    ensure_vendor_files()
    engine = ChatEngine()
    webview.create_window('AI Studio Pro', INDEX_FILE, js_api=engine, width=1250, height=850, background_color='#1c1c1e')
    webview.start(http_server=True, private_mode=False, storage_path=app_cache_dir, debug=False)