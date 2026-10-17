import re
import io
import html as html_lib
try:
    import docx
except ImportError:
    docx = None

# ---------- document export helpers (Word / HTML) ----------
try:
    from docx.oxml import OxmlElement
    from docx.shared import Pt
except Exception:
    OxmlElement = None
    Pt = None
from html.parser import HTMLParser

_RTL_RE = re.compile(r"[\u0590-\u08FF\uFB1D-\uFDFF\uFE70-\uFEFF]")   # Arabic / Hebrew letters
_ID_RE = re.compile(r"^[\w.-]{1,80}$")
_STATUS_PREFIX_RE = re.compile(r"^(?:>\s*(?:📂|🌐|⚠️|ℹ️)[^\n]*\n+)+")


class _DocxBuilder(HTMLParser):
    """Converts the (already cleaned) HTML of a reply into a Word document."""
    HEADINGS = {"h1": 1, "h2": 2, "h3": 3, "h4": 4, "h5": 4, "h6": 4}

    def __init__(self, document):
        super().__init__(convert_charrefs=True)
        self.doc = document
        self.cur = []
        self.bold = self.italic = self.code = self.quote = 0
        self.lists = []
        self.in_pre = False
        self.pre_buf = []
        self.rows = None
        self.row = None
        self.cell = None
        self.cell_header = False

    def _seg(self, text):
        return (text, self.bold > 0, self.italic > 0, self.code > 0)

    def _target(self):
        return self.cell if self.cell is not None else self.cur

    def _list_style(self):
        if not self.lists: return None
        depth = min(len(self.lists), 3)
        base = "List Bullet" if self.lists[-1] == "ul" else "List Number"
        return base if depth == 1 else f"{base} {depth}"

    def _write(self, paragraph, segments):
        for text, b, i, c in segments:
            run = paragraph.add_run(text)
            run.bold = True if b else None
            run.italic = True if i else None
            if c: run.font.name = "Courier New"
            if OxmlElement is not None and _RTL_RE.search(text):
                run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))
        if OxmlElement is not None and _RTL_RE.search("".join(s[0] for s in segments)):
            paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))

    @staticmethod
    def _trim(segs):
        segs = list(segs)
        if segs: segs[0] = (segs[0][0].lstrip(),) + tuple(segs[0][1:])
        if segs: segs[-1] = (segs[-1][0].rstrip(),) + tuple(segs[-1][1:])
        return [s for s in segs if s[0]]

    def flush(self, style=None):
        segs = self._trim(self.cur)
        self.cur = []
        if not segs: return
        try: p = self.doc.add_paragraph(style=style) if style else self.doc.add_paragraph()
        except KeyError: p = self.doc.add_paragraph()
        self._write(p, segs)

    def handle_starttag(self, tag, attrs):
        if tag in self.HEADINGS:
            self.flush(self._list_style())
        elif tag == "p":
            if self.cell is None and not self.lists: self.flush()
        elif tag in ("ul", "ol"):
            self.flush(self._list_style()); self.lists.append(tag)
        elif tag == "li":
            self.flush(self._list_style())
        elif tag in ("strong", "b"): self.bold += 1
        elif tag in ("em", "i"): self.italic += 1
        elif tag == "code":
            if not self.in_pre: self.code += 1
        elif tag == "pre":
            self.flush(); self.in_pre = True; self.pre_buf = []
        elif tag == "br":
            if self.in_pre: self.pre_buf.append("\n")
            else: self._target().append(self._seg("\n"))
        elif tag == "hr":
            self.flush(); self.doc.add_paragraph("-" * 60)
        elif tag == "blockquote":
            self.flush(); self.quote += 1
        elif tag == "table":
            self.flush(); self.rows = []
        elif tag == "tr":
            self.row = []
        elif tag in ("td", "th"):
            self.cell = []; self.cell_header = (tag == "th")

    def handle_endtag(self, tag):
        if tag in self.HEADINGS:
            self.flush(f"Heading {self.HEADINGS[tag]}")
        elif tag == "p":
            if self.cell is not None: self.cell.append(self._seg("\n"))
            elif self.lists: self.flush(self._list_style())
            elif self.quote: self.flush("Quote")
            else: self.flush()
        elif tag == "li":
            self.flush(self._list_style())
        elif tag in ("ul", "ol"):
            self.flush(self._list_style())
            if self.lists: self.lists.pop()
        elif tag in ("strong", "b"): self.bold = max(0, self.bold - 1)
        elif tag in ("em", "i"): self.italic = max(0, self.italic - 1)
        elif tag == "code":
            if not self.in_pre: self.code = max(0, self.code - 1)
        elif tag == "pre":
            text = "".join(self.pre_buf).rstrip("\n")
            p = self.doc.add_paragraph()
            run = p.add_run(text)
            run.font.name = "Courier New"
            if Pt is not None: run.font.size = Pt(9)
            self.in_pre = False
        elif tag == "blockquote":
            self.flush("Quote"); self.quote = max(0, self.quote - 1)
        elif tag in ("td", "th"):
            if self.row is not None and self.cell is not None: self.row.append((self.cell, self.cell_header))
            self.cell = None
        elif tag == "tr":
            if self.rows is not None and self.row is not None: self.rows.append(self.row)
            self.row = None
        elif tag == "table":
            self._build_table(); self.rows = None

    def handle_data(self, data):
        if self.in_pre:
            self.pre_buf.append(data); return
        data = re.sub(r"\s+", " ", data)
        if not data: return
        target = self._target()
        if not target and not data.strip(): return
        target.append(self._seg(data))

    def _build_table(self):
        rows = [r for r in (self.rows or []) if r]
        if not rows: return
        ncols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        try: table.style = "Table Grid"
        except Exception: pass
        for ri, row in enumerate(rows):
            for ci, (segs, header) in enumerate(row):
                para = table.cell(ri, ci).paragraphs[0]
                segs = self._trim(segs)
                self._write(para, [(t, b or header, i, c) for t, b, i, c in segs])
        self.doc.add_paragraph()


def html_to_docx_bytes(html_body):
    if docx is None or OxmlElement is None:
        raise RuntimeError("Word export needs the python-docx package (pip install python-docx).")
    document = docx.Document()
    builder = _DocxBuilder(document)
    builder.feed(html_body)
    builder.close()
    builder.flush()
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def build_html_document(title, body, autoprint=False):
    body = re.sub(r"<script\b.*?</script>", "", body or "", flags=re.S | re.I)
    css = """
      @page { margin: 14mm; }
      body { font-family: 'Segoe UI', Arial, sans-serif; color: #111; background: #fff; max-width: 920px; margin: 24px auto; padding: 0 18px; line-height: 1.6; font-size: 15px;
             -webkit-print-color-adjust: exact; print-color-adjust: exact; }
      h1, h2, h3, h4 { line-height: 1.25; margin: 1.2em 0 .5em; }
      h1 { font-size: 1.7em; } h2 { font-size: 1.4em; } h3 { font-size: 1.15em; }
      p, li, td, th, h1, h2, h3, h4, blockquote { unicode-bidi: plaintext; }
      table { border-collapse: collapse; margin: 1em 0; width: 100%; }
      th, td { border: 1px solid #bbb; padding: 6px 9px; text-align: start; vertical-align: top; }
      th { background: #e9edf2; }
      tr { page-break-inside: avoid; } thead { display: table-header-group; }
      pre { background: #f4f4f5; border: 1px solid #ddd; border-radius: 6px; padding: 10px; overflow-x: auto; white-space: pre-wrap; font-size: 13px; }
      code { font-family: Consolas, 'Courier New', monospace; background: #f1f1f2; padding: 1px 4px; border-radius: 4px; }
      pre code { background: none; padding: 0; }
      blockquote { border-inline-start: 4px solid #3b82f6; margin: 1em 0; padding: .2em 1em; color: #444; }
      hr { border: 0; border-top: 1px solid #ccc; margin: 1.2em 0; }
      img { max-width: 100%; }
      .noprint { background: #eef5ff; border: 1px solid #b6d4fe; padding: 10px 14px; border-radius: 8px; margin-bottom: 18px; font-size: 13px; }
      @media print { .noprint { display: none; } body { margin: 0; max-width: none; } }
    """
    hint = ('<div class="noprint">The print window opens automatically. Choose <b>Save as PDF</b> as the printer, then press Save. '
            'If it did not open, press <b>Ctrl + P</b>.</div>') if autoprint else ""
    script = "<script>window.addEventListener('load',function(){setTimeout(function(){window.print();},500);});</script>" if autoprint else ""
    return (f"<!DOCTYPE html><html><head><meta charset=\"utf-8\"><title>{html_lib.escape(title)}</title><style>{css}</style></head>"
            f"<body>{hint}{body}{script}</body></html>")



