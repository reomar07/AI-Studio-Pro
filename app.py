import os
import sys
import time
import string
import json
import base64
import threading
import webview
import pymupdf
import subprocess
import tempfile
import webbrowser
import textwrap
import copy
import shutil
from openai import OpenAI
from rank_bm25 import BM25Okapi
from ddgs import DDGS

try:
    import docx
except ImportError:
    docx = None

if getattr(sys, 'frozen', False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

CONFIG_FILE = os.path.join(BASE_DIR, "universal_config.json")
HISTORY_FILE = os.path.join(BASE_DIR, "chat_history.json")
SETTINGS_FILE = os.path.join(BASE_DIR, "chat_settings.json")
METADATA_FILE = os.path.join(BASE_DIR, "chat_metadata.json")
PROMPTS_FILE = os.path.join(BASE_DIR, "prompts.json")
INDEX_FILE = os.path.join(BASE_DIR, "index.html")

DEFAULT_PROFILES = {
    "active_profile_id": "default",
    "profiles": {
        "default": {"name": "OpenRouter", "base_url": "https://openrouter.ai/api/v1", "api_key": "", "model_name": "google/gemini-2.0-flash-001", "image_model": "dall-e-3", "custom_headers": ""},
        "local_gguf": {"name": "Local (LM Studio)", "base_url": "http://localhost:1234/v1", "api_key": "lm-studio", "model_name": "local-model", "image_model": "", "custom_headers": ""}
    }
}

class ChatEngine:
    def __init__(self):
        self.file_lock = threading.Lock()
        self.stream_lock = threading.Lock()
        self.stream_chunk = ""
        self.stream_is_done = False
        self.stream_error = None
        
        self.config = self._load_and_migrate_config()
        self.history = self._load_json(HISTORY_FILE, {})
        self._migrate_history_to_trees()
        self.chat_settings = self._load_json(SETTINGS_FILE, {})
        self.metadata = self._load_json(METADATA_FILE, {})
        self.prompts = self._load_json(PROMPTS_FILE, [])
        
        self.current_chat_id = f"chat_{int(time.time() * 1000)}"
        if self.current_chat_id not in self.history:
            self.history[self.current_chat_id] = {"nodes": {}, "current_node": None}
            
        self.is_aborted = False
        self.active_stream = None
        self.workspace_name = None
        self.workspace_docs = []
        self.workspace_meta = []
        self.bm25_engine = None
        self.session_files = {}

    def poll_stream(self):
        with self.stream_lock:
            chunk = self.stream_chunk
            self.stream_chunk = ""
            done = self.stream_is_done
            error = self.stream_error
            if done:
                self.stream_error = None
                self.stream_is_done = False
            return {"chunk": chunk, "done": done, "error": error}

    def open_external_url(self, url):
        if url and (url.startswith("http://") or url.startswith("https://")):
            webbrowser.open(url)
            return True
        return False

    def save_code_to_file(self, code_string):
        try:
            dialog_type = getattr(webview.FileDialog, 'SAVE', getattr(webview, 'SAVE_DIALOG', 1))
            result = webview.windows[0].create_file_dialog(dialog_type, save_filename='snippet.txt')
            if not result or len(result) == 0: return None
            
            with open(result[0], 'w', encoding='utf-8') as f:
                f.write(code_string)
            webview.windows[0].evaluate_js("showToast('File saved successfully!')")
            return True
        except Exception as e:
            safe_err = str(e).replace("'", "\\'")[:80]
            webview.windows[0].evaluate_js(f"showToast('Save failed: {safe_err}')")
            return False

    def export_chat_dialog(self, chat_id, export_format="md"):
        try:
            messages = self._get_linear_branch(self.history.get(chat_id))
            if not messages: 
                webview.windows[0].evaluate_js("showToast('No messages to export.')")
                return False
            
            content = ""
            if export_format == "json":
                content = json.dumps(messages, indent=2)
            else:
                content = f"# Chat Export: {messages[0].get('title', 'Conversation')}\n\n"
                for m in messages:
                    role_name = "**User**" if m['role'] == 'user' else "**Assistant**"
                    text = m.get('display_text', m.get('content', ''))
                    if isinstance(text, list): 
                        text = "\n".join([item.get('text', '') for item in text if item.get('type') == 'text'])
                    content += f"{role_name}:\n{text}\n\n---\n\n"
            
            dialog_type = getattr(webview.FileDialog, 'SAVE', getattr(webview, 'SAVE_DIALOG', 1))
            safe_title = "".join([c for c in messages[0].get('title', 'Conversation') if c.isalpha() or c.isdigit() or c==' ']).rstrip()
            safe_title = safe_title.replace(" ", "_")
            if not safe_title: safe_title = "Conversation"
            
            result = webview.windows[0].create_file_dialog(
                dialog_type, 
                save_filename=f"{safe_title}.{export_format}"
            )
            
            if result and len(result) > 0:
                with open(result[0], 'w', encoding='utf-8') as f:
                    f.write(content)
                safe_path = str(result[0]).replace("\\", "\\\\").replace("'", "\\'")
                webview.windows[0].evaluate_js(f"showToast('Success! Saved to: {safe_path}')")
                return True
            else:
                webview.windows[0].evaluate_js("showToast('Export cancelled.')")
                return False
        except Exception as e:
            safe_err = str(e).replace("'", "\\'")[:80]
            webview.windows[0].evaluate_js(f"showToast('Export failed: {safe_err}')")
            return False

    def _tokenize(self, text):
        return text.lower().translate(str.maketrans('', '', string.punctuation)).split()

    def _load_json(self, path, default):
        if os.path.exists(path):
            try:
                with open(path, "r", encoding="utf-8") as f: return json.load(f)
            except Exception: pass
        return default

    def _save_json(self, path, data):
        with self.file_lock:
            temp_path = f"{path}.tmp_{os.urandom(2).hex()}"
            try:
                with open(temp_path, "w", encoding="utf-8") as f:
                    json.dump(data, f, indent=4)
                os.replace(temp_path, path)
            except Exception as e:
                print(f"Error saving {path}: {e}")
                if os.path.exists(temp_path):
                    try: os.remove(temp_path)
                    except: pass

    def _load_and_migrate_config(self):
        data = self._load_json(CONFIG_FILE, None)
        if not data or "profiles" not in data:
            self._save_json(CONFIG_FILE, DEFAULT_PROFILES)
            return DEFAULT_PROFILES
        return data

    def _migrate_history_to_trees(self):
        dirty = False
        for cid, chat_data in list(self.history.items()):
            if isinstance(chat_data, list):
                dirty = True
                new_chat = {"nodes": {}, "current_node": None}
                prev_id = None
                for m in chat_data:
                    nid = "msg_" + os.urandom(4).hex()
                    m["id"] = nid
                    m["parent"] = prev_id
                    m["children"] = []
                    new_chat["nodes"][nid] = m
                    if prev_id: new_chat["nodes"][prev_id]["children"].append(nid)
                    prev_id = nid
                new_chat["current_node"] = prev_id
                self.history[cid] = new_chat
        if dirty: self._save_json(HISTORY_FILE, self.history)

    def _get_linear_branch(self, chat_data):
        if not chat_data or not chat_data.get("current_node"): return []
        nodes = chat_data["nodes"]
        curr = chat_data["current_node"]
        path = []
        while curr and curr in nodes:
            path.append(curr)
            curr = nodes[curr]["parent"]
        path.reverse()
        linear = []
        for nid in path:
            msg = dict(nodes[nid])
            pid = msg["parent"]
            if pid and pid in nodes:
                siblings = nodes[pid]["children"]
                msg["branch_count"] = len(siblings)
                msg["branch_index"] = siblings.index(nid) + 1
            else:
                msg["branch_count"] = 1
                msg["branch_index"] = 1
            linear.append(msg)
        return linear

    def _add_node(self, chat_id, role, content, display_text, title, is_image_mode, files=None, image=None):
        chat = self.history[chat_id]
        nid = "msg_" + os.urandom(4).hex()
        pid = chat.get("current_node")
        msg = {
            "id": nid, "parent": pid, "children": [], "role": role, 
            "content": content, "display_text": display_text, "title": title, "is_image_mode": is_image_mode,
            "files": files or [], "image": image
        }
        chat["nodes"][nid] = msg
        if pid and pid in chat["nodes"]: chat["nodes"][pid]["children"].append(nid)
        chat["current_node"] = nid
        return nid

    def switch_branch(self, chat_id, msg_id, direction):
        chat = self.history[chat_id]
        if msg_id not in chat["nodes"]: return False
        msg = chat["nodes"][msg_id]
        pid = msg["parent"]
        if not pid: return False
        siblings = chat["nodes"][pid]["children"]
        idx = siblings.index(msg_id)
        new_idx = (idx + direction) % len(siblings)
        curr = siblings[new_idx]
        while chat["nodes"][curr]["children"]:
            curr = chat["nodes"][curr]["children"][-1]
        chat["current_node"] = curr
        self._save_json(HISTORY_FILE, self.history)
        return True

    def get_config(self): return self.config
    def get_history(self): return {cid: self._get_linear_branch(cdata) for cid, cdata in self.history.items()}
    def get_metadata(self): return self.metadata
    def get_chat_settings(self, chat_id): return self.chat_settings.get(chat_id, {})
    def get_prompts(self): return self.prompts

    def save_prompt(self, title, content, prompt_id=None):
        if prompt_id:
            for p in self.prompts:
                if p["id"] == prompt_id:
                    p["title"], p["content"] = title, content
                    break
        else:
            self.prompts.append({"id": f"p_{os.urandom(4).hex()}", "title": title, "content": content})
        self._save_json(PROMPTS_FILE, self.prompts)
        return self.prompts

    def delete_prompt(self, prompt_id):
        self.prompts = [p for p in self.prompts if p["id"] != prompt_id]
        self._save_json(PROMPTS_FILE, self.prompts)
        return self.prompts

    def save_config(self, cfg):
        self.config.update(cfg)
        active_id = self.config.get("active_profile_id", "default")
        profile = self.config.get("profiles", {}).get(active_id, {})
        if profile:
            self.config["api_key"] = profile.get("api_key", "")
            self.config["base_url"] = profile.get("base_url", "")
            self.config["model_name"] = profile.get("model_name", "")
            self.config["image_model"] = profile.get("image_model", "dall-e-3")
        self._save_json(CONFIG_FILE, self.config)
        return True

    def start_new_chat(self, is_image_mode=False):
        self.current_chat_id = f"chat_{int(time.time() * 1000)}"
        self.history[self.current_chat_id] = {"nodes": {}, "current_node": None}
        self.session_files.clear()
        return self.current_chat_id

    def set_current_chat(self, chat_id): 
        self.current_chat_id = chat_id
        if self.current_chat_id not in self.history:
            self.history[self.current_chat_id] = {"nodes": {}, "current_node": None}

    def toggle_pin_chat(self, chat_id):
        if chat_id not in self.metadata: self.metadata[chat_id] = {}
        self.metadata[chat_id]["pinned"] = not self.metadata[chat_id].get("pinned", False)
        self._save_json(METADATA_FILE, self.metadata)
        return self.metadata[chat_id]["pinned"]

    def delete_chat(self, chat_id):
        for store, file_path in [(self.history, HISTORY_FILE), (self.chat_settings, SETTINGS_FILE), (self.metadata, METADATA_FILE)]:
            if chat_id in store:
                del store[chat_id]
                self._save_json(file_path, store)
        return True

    def clear_all_chats(self):
        self.history, self.chat_settings, self.metadata = {}, {}, {}
        self._save_json(HISTORY_FILE, self.history); self._save_json(SETTINGS_FILE, self.chat_settings); self._save_json(METADATA_FILE, self.metadata)
        self.session_files.clear()
        return True

    def stop_generation(self):
        self.is_aborted = True
        if self.active_stream:
            try: self.active_stream.close()
            except: pass

    def rewind_chat(self, chat_id, msg_id):
        chat = self.history.get(chat_id)
        if chat and msg_id in chat["nodes"]:
            chat["current_node"] = chat["nodes"][msg_id]["parent"]
            self._save_json(HISTORY_FILE, self.history)
        return True

    def _get_python_interpreter(self):
        if not getattr(sys, 'frozen', False):
            return sys.executable
        for cmd in ['python', 'py', 'python3']:
            found = shutil.which(cmd)
            if found:
                if os.path.abspath(found).lower() != os.path.abspath(sys.executable).lower():
                    return found
        possible_dirs = [
            os.path.expandvars(r"%LOCALAPPDATA%\Programs\Python"),
            r"C:\Python312", r"C:\Python311", r"C:\Python310", r"C:\Python39"
        ]
        for base in possible_dirs:
            if os.path.exists(base):
                for root, dirs, files in os.walk(base):
                    if "python.exe" in files:
                        p_path = os.path.join(root, "python.exe")
                        if "scripts" not in p_path.lower():
                            return p_path
        return None

    def execute_local_code(self, code_string):
        python_bin = self._get_python_interpreter()
        if not python_bin:
            return "Execution Error: Python is not installed or not in your Windows PATH. Please install Python to run scripts locally."
        
        temp_path = None
        try:
            with tempfile.NamedTemporaryFile(mode='w', suffix='.py', delete=False, encoding='utf-8') as f:
                f.write(code_string)
                temp_path = f.name
            
            creation_flags = 0
            if sys.platform == 'win32':
                creation_flags = subprocess.CREATE_NO_WINDOW

            result = subprocess.run(
                [python_bin, temp_path],
                capture_output=True,
                text=True,
                timeout=15,
                creationflags=creation_flags
            )
            output = result.stdout
            if result.stderr:
                output += f"\nErrors:\n{result.stderr}"
            return output.strip() or "Code executed successfully with no console output."
        except subprocess.TimeoutExpired:
            return "Execution Error: Script timed out after 15 seconds."
        except Exception as e:
            return f"Execution Error: {str(e)}"
        finally:
            if temp_path and os.path.exists(temp_path):
                try: os.remove(temp_path)
                except: pass

    def pick_file(self):
        file_types = ('All files (*.*)',)
        try:
            dialog_type = getattr(webview.FileDialog, 'OPEN', getattr(webview, 'OPEN_DIALOG', 0))
            result = webview.windows[0].create_file_dialog(dialog_type, allow_multiple=False, file_types=file_types)
            if not result or len(result) == 0: return None
            path = result[0]
            ext = path.lower().split('.')[-1]
            name = os.path.basename(path)
            
            file_id = os.urandom(8).hex()

            if ext in ['png', 'jpg', 'jpeg', 'webp', 'gif', 'bmp']:
                with open(path, "rb") as f: b64 = base64.b64encode(f.read()).decode('utf-8')
                data_url = f"data:image/{ext};base64,{b64}"
                self.session_files[file_id] = {"name": name, "is_image": True, "data_url": data_url}
                return {"id": file_id, "name": name, "is_image": True, "data_url": data_url}

            raw_bytes = open(path, "rb").read()
            text_data = ""
            if ext == 'pdf': text_data = "\n".join(p.get_text() for p in pymupdf.open(stream=raw_bytes, filetype="pdf"))
            elif ext == 'docx' and docx:
                import io
                text_data = "\n".join(p.text for p in docx.Document(io.BytesIO(raw_bytes)).paragraphs)
            else:
                text_data = raw_bytes.decode('utf-8', errors='ignore')
                
            self.session_files[file_id] = {"name": name, "is_image": False, "text": text_data}
            return {"id": file_id, "name": name, "is_image": False}
        except Exception: return None

    def attach_workspace(self):
        try:
            dialog_type = getattr(webview.FileDialog, 'FOLDER', getattr(webview, 'FOLDER_DIALOG', 2))
            result = webview.windows[0].create_file_dialog(dialog_type)
            if not result or len(result) == 0: return None
            folder_path = result[0]
            safe_name = os.path.basename(folder_path).replace("'", "\\'")
            webview.windows[0].evaluate_js(f"showToast('Scanning {safe_name}...')")
            threading.Thread(target=self._index_folder, args=(folder_path,), daemon=True).start()
            return os.path.basename(folder_path)
        except Exception:
            return None

    def _index_folder(self, folder_path):
        try:
            self.workspace_docs = []
            self.workspace_meta = []
            self.bm25_engine = None
            self.workspace_name = os.path.basename(folder_path)
            
            exclude_dirs = {'.gradle', '.idea', 'build', 'node_modules', '.git', 'intermediates', '__pycache__', '.venv', 'env'}
            
            for root, dirs, files in os.walk(folder_path):
                dirs[:] = [d for d in dirs if d not in exclude_dirs]
                for file in files:
                    ext = file.split('.')[-1].lower()
                    path = os.path.join(root, file)
                    text = ""
                    if ext in ['kt', 'py', 'md', 'txt', 'java', 'xml', 'json', 'csv', 'html', 'js']:
                        try:
                            with open(path, 'r', encoding='utf-8') as f:
                                text = f.read()
                        except: pass
                    elif ext == 'pdf':
                        try: text = "\n".join(page.get_text() for page in pymupdf.open(path))
                        except: pass
                    
                    if text:
                        chunks = textwrap.wrap(text, width=1200, break_long_words=False, replace_whitespace=False)
                        for chunk in chunks:
                            if chunk.strip():
                                self.workspace_docs.append(chunk)
                                self.workspace_meta.append({"filename": file})
                        
            if self.workspace_docs:
                tokenized_corpus = [self._tokenize(doc) for doc in self.workspace_docs]
                self.bm25_engine = BM25Okapi(tokenized_corpus)
                webview.windows[0].evaluate_js(f"showToast('Indexed {len(self.workspace_docs)} chunks successfully.')")
            else:
                webview.windows[0].evaluate_js("showToast('No readable text files found.')")
        except Exception as e:
            safe_err = str(e).replace("'", "\\'")[:80]
            webview.windows[0].evaluate_js(f"showToast('Indexing error: {safe_err}')")

    def process_frontend_file(self, file_obj):
        file_id = file_obj.get("id")
        if file_id and file_id in self.session_files:
            return self.session_files[file_id]
        
        try:
            name, data_url = file_obj.get("name", "unnamed"), file_obj.get("data", "")
            ext = name.split('.')[-1].lower()
            if not data_url or "," not in data_url: return None
            b64_data = data_url.split(",")[1]

            if ext in ['png', 'jpg', 'jpeg', 'webp', 'gif', 'bmp']: return {"name": name, "is_image": True, "data_url": data_url}
            raw_bytes = base64.b64decode(b64_data)
            if ext == 'pdf': return {"name": name, "is_image": False, "text": "\n".join(p.get_text() for p in pymupdf.open(stream=raw_bytes, filetype="pdf"))}
            if ext == 'docx' and docx:
                import io
                return {"name": name, "is_image": False, "text": "\n".join(p.text for p in docx.Document(io.BytesIO(raw_bytes)).paragraphs)}
            return {"name": name, "is_image": False, "text": raw_bytes.decode('utf-8', errors='ignore')}
        except Exception: return None

    def _generate_title_thread(self, chat_id, prompt_text, profile):
        try:
            custom_headers_raw = profile.get("custom_headers", "").strip()
            headers = json.loads(custom_headers_raw) if custom_headers_raw.startswith("{") else None
            
            client = OpenAI(base_url=profile.get("base_url") or None, api_key=profile.get("api_key"), default_headers=headers, timeout=120.0)
            response = client.chat.completions.create(
                model=profile.get("model_name", "gpt-4o"),
                messages=[{"role": "system", "content": "Summarize this user request in 2 to 4 words. Respond ONLY with the title. No quotes, no punctuation."}, {"role": "user", "content": str(prompt_text)[:500]}],
                max_tokens=10, temperature=0.3
            )
            chat = self.history.get(chat_id)
            if chat and chat.get("current_node"):
                curr = chat["current_node"]
                while chat["nodes"][curr]["parent"]: curr = chat["nodes"][curr]["parent"]
                chat["nodes"][curr]["title"] = response.choices[0].message.content.strip().replace('"', '')
                self._save_json(HISTORY_FILE, self.history)
        except Exception: pass

    def send_prompt_stream(self, user_prompt, file_data_list, params):
        self.is_aborted = False
        threading.Thread(target=self._stream_worker, args=(user_prompt, file_data_list, params), daemon=True).start()

    def _stream_worker(self, user_prompt, file_data_list, params):
        with self.stream_lock:
            self.stream_chunk = ""
            self.stream_is_done = False
            self.stream_error = None

        sources_list = []
        target_chat = self.current_chat_id 
        if target_chat not in self.history:
            self.history[target_chat] = {"nodes": {}, "current_node": None}
            
        self.chat_settings[target_chat] = params
        self._save_json(SETTINGS_FILE, self.chat_settings)

        active_id = self.config.get("active_profile_id", "default")
        profile = self.config.get("profiles", {}).get(active_id, {})
        
        api_key = profile.get("api_key") or self.config.get("api_key", "")
        base_url = profile.get("base_url") or self.config.get("base_url", "")
        model_name = profile.get("model_name") or self.config.get("model_name", "google/gemini-2.0-flash-001")
        
        custom_headers_raw = profile.get("custom_headers", "").strip()
        custom_headers = None
        if custom_headers_raw.startswith("{") and custom_headers_raw.endswith("}"):
            try: custom_headers = json.loads(custom_headers_raw)
            except: pass

        is_image_gen = params.get("is_image_mode", False)
        is_regenerate = params.get("regenerate", False)
        active_workspace = params.get("active_workspace")

        if not api_key and "localhost" not in str(base_url):
            self._add_node(target_chat, "user", user_prompt, user_prompt, (user_prompt or "Chat")[:28], False)
            err_msg = "**Configuration Warning**: No API Key detected. Please open Settings and configure your Provider Profile."
            self._add_node(target_chat, "assistant", err_msg, err_msg, "", False)
            self._save_json(HISTORY_FILE, self.history)
            with self.stream_lock:
                self.stream_error = err_msg
                self.stream_is_done = True
            return

        current_turn_payload = None
        if not is_regenerate:
            doc_texts, image_urls, doc_meta_list = [], [], []
            
            if file_data_list:
                for f in file_data_list:
                    processed = self.process_frontend_file(f)
                    if not processed: continue
                    if processed["is_image"]: image_urls.append(processed["data_url"])
                    else: 
                        doc_texts.append(f"--- Document: {processed['name']} ---\n{processed['text']}")
                        doc_meta_list.append({"name": processed["name"]})

            workspace_context = ""
            if active_workspace and self.bm25_engine and user_prompt and user_prompt.strip():
                try:
                    tokenized_query = self._tokenize(user_prompt)
                    scores = self.bm25_engine.get_scores(tokenized_query)
                    top_n_indices = sorted(range(len(scores)), key=lambda i: scores[i], reverse=True)[:6]
                    if sum(scores) > 0:
                        workspace_context += f"\n[Local Workspace Files: {active_workspace}]\n"
                        for idx in top_n_indices:
                            doc = self.workspace_docs[idx]
                            meta = self.workspace_meta[idx]
                            workspace_context += f"--- File: {meta['filename']} ---\n{doc}\n\n"
                except: pass

            is_web_search = params.get("web_search", False)
            web_context = ""
            if is_web_search and user_prompt and user_prompt.strip():
                try:
                    with self.stream_lock:
                        self.stream_chunk += "> 🌐 *Searching live web...*\n\n"
                    results = list(DDGS().text(user_prompt, max_results=4))
                    if results:
                        web_context += "\n[Live Web Search Reference Material]\n"
                        for res in results:
                            title = str(res.get('title', 'Web Result')).replace('"', '&quot;')
                            url = str(res.get('href', ''))
                            body = str(res.get('body', ''))
                            if url.startswith("http://") or url.startswith("https://"):
                                sources_list.append({"title": title, "url": url})
                                web_context += f"Source: {title} ({url})\nSnippet: {body}\n\n"
                    else:
                        web_context += "\n[Live Web Search: No matching results found online.]\n"
                except Exception as e:
                    pass

            if image_urls or doc_texts or workspace_context or web_context:
                combined_text = ""
                if workspace_context: combined_text += f"{workspace_context}\n"
                if doc_texts: combined_text += f"{''.join(doc_texts)}\n"
                if web_context: combined_text += f"{web_context}\n"
                combined_text += f"User Request: {user_prompt}"
                
                current_turn_payload = [{"type": "text", "text": combined_text}]
                for url in image_urls: current_turn_payload.append({"type": "image_url", "image_url": {"url": url}})
            else:
                current_turn_payload = user_prompt

            clean_history_content = [{"type": "text", "text": str(user_prompt or "")}] if image_urls else user_prompt
            if image_urls:
                for url in image_urls: clean_history_content.append({"type": "image_url", "image_url": {"url": url}})

            display_title = (user_prompt or "Attachment")[:28]
            display_label = user_prompt or f"[{len(file_data_list or [])} Files Attached]"
            primary_image = image_urls[0] if image_urls else None

            self._add_node(target_chat, "user", clean_history_content, display_label, display_title, is_image_gen, files=doc_meta_list, image=primary_image)

            if len(self._get_linear_branch(self.history[target_chat])) == 1 and not is_image_gen:
                threading.Thread(target=self._generate_title_thread, args=(target_chat, user_prompt, {"api_key": api_key, "base_url": base_url, "model_name": model_name}), daemon=True).start()

        linear_history = self._get_linear_branch(self.history[target_chat])
        client = OpenAI(base_url=base_url or None, api_key=api_key, default_headers=custom_headers, timeout=120.0)
        assistant_node_id = None

        if is_image_gen:
            try:
                img_prompt = linear_history[-1].get("display_text", "") if linear_history else ""
                image_model = profile.get("image_model") or "dall-e-3"
                response = client.images.generate(model=image_model, prompt=img_prompt, size="1024x1024", quality="standard", n=1)
                img_url = response.data[0].url
                output_md = f"![Generated Image]({img_url})\n\n[Open Full Size Image]({img_url})"
                self._add_node(target_chat, "assistant", output_md, output_md, "", False, image=img_url)
                self._save_json(HISTORY_FILE, self.history)
                with self.stream_lock:
                    self.stream_chunk += output_md
                    self.stream_is_done = True
            except Exception as e: 
                err = f"**Image Generation Notice**: {str(e)}"
                self._add_node(target_chat, "assistant", err, err, "", False)
                self._save_json(HISTORY_FILE, self.history)
                with self.stream_lock:
                    self.stream_error = err
                    self.stream_is_done = True
            return

        try:
            limit = params.get("context_limit", 10)
            messages_payload = [{"role": "system", "content": params["system_prompt"]}] if params.get("system_prompt") else []
            
            for m in linear_history[-limit:-1]:
                content = m.get("content", "")
                if isinstance(content, str) and len(content) > 3500:
                    content = content[:3500] + "\n...[Historical turn truncated]..."
                elif isinstance(content, list):
                    content = copy.deepcopy(content)
                    for item in content:
                        if item.get("type") == "text" and isinstance(item.get("text"), str) and len(item["text"]) > 3500:
                            item["text"] = item["text"][:3500] + "\n...[Historical turn truncated]..."
                messages_payload.append({"role": m["role"], "content": content})
            
            if linear_history:
                last_turn = linear_history[-1]
                active_content = current_turn_payload if (current_turn_payload and not is_regenerate) else last_turn.get("content", "")
                messages_payload.append({"role": last_turn["role"], "content": active_content})

            self.active_stream = client.chat.completions.create(model=model_name, messages=messages_payload, temperature=params.get("temperature", 0.7), stream=True)

            assistant_node_id = self._add_node(target_chat, "assistant", "", "", "", False)
            full_reply = ""

            if active_workspace and not is_regenerate:
                status_msg = f"> 📂 *Scanned local workspace ({active_workspace}).*\n\n"
                full_reply += status_msg
                with self.stream_lock:
                    self.stream_chunk += status_msg
            
            for chunk in self.active_stream:
                if self.is_aborted:
                    self.active_stream.close()
                    cancel_msg = "\n\n*(Cancelled by user)*"
                    full_reply += cancel_msg
                    with self.stream_lock:
                        self.stream_chunk += cancel_msg
                    break
                if chunk.choices and chunk.choices[0].delta.content:
                    token = chunk.choices[0].delta.content
                    full_reply += token
                    with self.stream_lock:
                        self.stream_chunk += token

            self.history[target_chat]["nodes"][assistant_node_id]["content"] = full_reply
            self.history[target_chat]["nodes"][assistant_node_id]["display_text"] = full_reply
            self.history[target_chat]["nodes"][assistant_node_id]["sources"] = sources_list 
            self._save_json(HISTORY_FILE, self.history)
            
            with self.stream_lock:
                self.stream_is_done = True

        except Exception as e: 
            err = f"\n\n**API Error**: `{str(e)}`"
            if assistant_node_id:
                self.history[target_chat]["nodes"][assistant_node_id]["content"] = err
                self.history[target_chat]["nodes"][assistant_node_id]["display_text"] = err
            else:
                self._add_node(target_chat, "assistant", err, err, "", False)
            self._save_json(HISTORY_FILE, self.history)
            
            with self.stream_lock:
                self.stream_error = err
                self.stream_is_done = True


if __name__ == '__main__':
    import multiprocessing
    multiprocessing.freeze_support()
    app_cache_dir = os.path.join(BASE_DIR, 'webview_cache')
    os.makedirs(app_cache_dir, exist_ok=True)
    engine = ChatEngine()
    webview.create_window('AI Studio Pro', INDEX_FILE, js_api=engine, width=1250, height=850, background_color='#1c1c1e')
    webview.start(http_server=True, private_mode=False, storage_path=app_cache_dir, debug=False)